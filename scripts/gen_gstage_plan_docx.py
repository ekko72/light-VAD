# -*- coding: utf-8 -*-
"""生成《G 阶段执行规划：轻量神经 VAD》Word 文档，保存到桌面。"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = r"C:\Users\20547\Desktop\G阶段执行规划-轻量神经VAD-研一上2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)


def set_ea(run, ea="宋体"):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), ea)


def para(text, bold=False, size=10.5, align=None, ea="宋体"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    set_ea(r, ea)
    if align is not None:
        p.alignment = align
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    set_ea(r)
    return p


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.size = Pt(14 if level == 1 else 12)
        set_ea(r, "黑体")
        r.font.color.rgb = RGBColor(0, 0, 0)
    return h


# ===== 标题 =====
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("G 阶段执行规划：轻量神经 VAD")
r.bold = True
r.font.size = Pt(18)
set_ea(r, "黑体")

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("研一上学期（2026.9–2026.12）· 每周投入约 10–15h · 路线：G 垫脚石 → A' 主线 → D 逃生舱")
r.font.size = Pt(10.5)
set_ea(r, "楷体")

# ===== 一、定位与目标 =====
heading("一、定位与目标", 1)
bullet("为什么先做 VAD：问题小、数据公开（DNS 数据 + 噪声合成）、评测指标明确（AUC/EER/延迟），2–3 个月能完整闭环；")
bullet("VAD 是 A'（自适应计算 SE）的门控零件——现在投入的时间后面全部复用，不是“先水一篇”；")
bullet("最终交付物：一个 <10K 参数、流式、INT8 可导出、多噪声鲁棒的神经 VAD；")
bullet("1 件专利交底；1 篇短文/中文核心投稿；简历上可写的项目条目。")

# ===== 二、里程碑周计划 =====
heading("二、里程碑周计划（验收驱动）", 1)

rows = [
    ("9月 W1-2", "真跑 D2L 的 CNN/RNN 章节（不是读）；装 PyTorch；下载数据子集",
     "训练循环能自己写出来"),
    ("9月 W3-4", "特征管线 + DataLoader；能量 VAD、WebRTC VAD 基线评测",
     "拿到基线分数表（后面所有对比的锚点）"),
    ("10月 W1-2", "第一个 GRU VAD 训练 + 评测闭环",
     "指标超过能量基线"),
    ("10月 W3-4", "结构瘦身 + 多 SNR/多噪声评测",
     "达到 <10K 参数，AUC ≥0.95（示例线）"),
    ("11月 W1-2", "流式化改造（状态传递、因果性验证）+ 延迟测量",
     "逐帧输出可用，延迟达标"),
    ("11月 W3-4", "ONNX + INT8 PTQ，量化前后指标对比；误报率调优",
     "量化后 AUC 掉 <1%（示例线）"),
    ("12月 W1-2", "专利交底 1–2 件；短文大纲 + 图表",
     "交底书给导师看"),
    ("12月 W3-4", "短文初稿；简历条目；README 整理",
     "简历能写条目；代码仓库干净可复现"),
]

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ("周次", "内容", "验收标准")
for i, htxt in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(htxt)
    r.bold = True
    r.font.size = Pt(10.5)
    set_ea(r, "黑体")

for week, content, accept in rows:
    cells = table.add_row().cells
    for i, txt in enumerate((week, content, accept)):
        cells[i].text = ""
        p = cells[i].paragraphs[0]
        r = p.add_run(txt)
        r.font.size = Pt(10)
        set_ea(r)

widths = (Cm(2.2), Cm(7.6), Cm(6.0))
for row in table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

# ===== 三、知识清单 =====
heading("三、知识清单", 1)
bullet("深度学习补课：D2L 前两章之后补 CNN/RNN（LSTM/GRU）、训练流程、loss 设计（BCE/Focal）、过拟合处理；")
bullet("音频侧：帧/帧移/窗、FBank/MFCC/短时能量特征、流式窗口设计；")
bullet("VAD 领域：传统能量/过零率（毕设底子，只当对比基线）、WebRTC VAD 当参照、神经 VAD 怎么在极低参数下保鲁棒；")
bullet("部署接口：流式推理（帧级输出 + 状态管理）、INT8 量化概念（PTQ/QAT）、ONNX 导出——为 RKNN/树莓派预留。")

# ===== 四、环境与数据 =====
heading("四、环境与数据（已就绪）", 1)
bullet("环境已预装（2026-08-13）：项目在 C:\\Users\\20547\\Desktop\\论文p12\\轻量VAD项目\\.venv，Python 3.12 + PyTorch 2.9.1（cu128 GPU，RTX 5060 已验证）；")
bullet("自检命令：python scripts\\verify_env.py（15 项全部通过）；")
bullet("数据下载脚本：scripts\\download_data.py（LibriSpeech dev-clean/test-clean + MUSAN，约 2.2GB，9 月再下）；DNS 数据集太大，先不装。")

# ===== 五、风险与收缩方案 =====
heading("五、风险与收缩方案", 1)
bullet("VAD 结果不够惊艳没关系——对比实验 + 工程化本身就能撑一篇；")
bullet("真被课程/横向挤压：砍到“专利交底 + 内部报告”即可，不影响后续；")
bullet("A' 的 v1 门控先用简单能量门控顶着（练手用，不进论文）。")

# ===== 六、每周执行节奏建议 =====
heading("六、每周执行节奏建议", 1)
bullet("固定闭环：学知识点 → 看参考代码 → 自己写 → AI 检验改错 → 记到 Obsidian；")
bullet("每周日花 30 分钟对照上表验收标准自检，偏差超过一周就启动收缩方案；")
bullet("代码、数据、笔记分开放：代码在项目文件夹，笔记在 Obsidian，数据脚本化可复现。")

para("", size=8)
para("生成日期：2026-08-13 · 配合 Obsidian「04_学习/第二阶段_神经VAD」使用", size=8)

doc.save(OUT)
print("saved:", OUT)
